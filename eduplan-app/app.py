import io
import json
import time
import streamlit as st
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from google import genai
from google.genai import types

# ----------------------------------------------------
# 1. การตั้งค่าระบบความปลอดภัยและส่วนแสดงผล (UI)
# ----------------------------------------------------
SYSTEM_PASSCODE = "0863449483"

st.set_page_config(
    page_title="EduPlan Pro | ระบบจัดทำโครงการสอนอัตโนมัติ",
    page_icon="🎓",
    layout="wide"
)

st.markdown("""
<style>
    @import url('https://fonts.googleapis.com/css2?family=Prompt:wght@300;400;500;600;700&display=swap');
    * { font-family: 'Prompt', sans-serif !important; }
    .hero-banner {
        background: linear-gradient(135deg, #1e3c72 0%, #2a5298 50%, #00c6ff 100%);
        border-radius: 18px; padding: 25px 30px; color: white; margin-bottom: 25px;
    }
    .hero-title { font-size: 2.1rem; font-weight: 700; margin-bottom: 6px; }
    .hero-desc { font-size: 1rem; opacity: 0.95; font-weight: 300; line-height: 1.5; }
    .box-header { font-size: 1.15rem; font-weight: 600; color: #1e3c72; margin-bottom: 10px; }
    .copyright-card {
        background: #f8fafc; border-left: 4px solid #1e3c72; border-radius: 8px; padding: 12px 16px; margin-top: 20px;
    }
</style>
""", unsafe_allow_html=True)

if "authenticated" not in st.session_state:
    st.session_state.authenticated = False

def login_gate():
    c1, c2, c3 = st.columns([1, 1.8, 1])
    with c2:
        st.markdown("<br><br>", unsafe_allow_html=True)
        st.markdown("""
        <div style="text-align: center; margin-bottom: 20px;">
            <div style="font-size: 3.5rem;">🏛️</div>
            <h2 style="font-weight: 700; color: #1e3c72;">EduPlan Pro (AI Engine)</h2>
            <p style="color: #64748b;">ระบบจัดทำโครงการสอนมาตรฐาน สอศ. ด้วย AI</p>
        </div>
        """, unsafe_allow_html=True)
        passcode = st.text_input("🔑 รหัสปลดล็อกสิทธิ์เข้าใช้งาน:", type="password")
        if st.button("🔓 ปลดล็อกและเข้าสู่ระบบ", type="primary", use_container_width=True):
            if passcode == SYSTEM_PASSCODE:
                st.session_state.authenticated = True
                st.rerun()
            else:
                st.error("❌ รหัสผ่านไม่ถูกต้อง")
        st.markdown("""
        <div class="copyright-card">
            <b>ลิขสิทธิ์และการพัฒนา</b><br>
            นายณัฐวุฒิ หล้าปงสาย ครูผู้ช่วย แผนกวิชาการจัดการโลจิสติกส์และซัพพลายเชน<br>
            วิทยาลัยเทคนิคจันทบุรี | All Rights Reserved © 2026
        </div>
        """, unsafe_allow_html=True)

if not st.session_state.authenticated:
    login_gate()
    st.stop()

# ----------------------------------------------------
# 2. ฟังก์ชันจัดการ Word XML และระบบเลขหน้า
# ----------------------------------------------------
def add_page_number_field(run):
    """แทรกฟิลด์ PAGE เพื่อให้ Word คำนวณเลขหน้าตามจริงของเอกสาร"""
    fldChar1 = OxmlElement('w:fldChar')
    fldChar1.set(qn('w:fldCharType'), 'begin')
    instrText = OxmlElement('w:instrText')
    instrText.set(qn('xml:space'), 'preserve')
    instrText.text = "PAGE"
    fldChar2 = OxmlElement('w:fldChar')
    fldChar2.set(qn('w:fldCharType'), 'separate')
    fldChar3 = OxmlElement('w:fldChar')
    fldChar3.set(qn('w:fldCharType'), 'end')
    
    r = run._r
    r.append(fldChar1)
    r.append(instrText)
    r.append(fldChar2)
    r.append(fldChar3)

def set_cell_font(cell, font_name="TH SarabunPSK", font_size=Pt(14), bold=False):
    for p in cell.paragraphs:
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.0
        for run in p.runs:
            run.font.name = font_name
            run.font.size = font_size
            run.font.bold = bold
            rPr = run._r.get_or_add_rPr()
            rFonts = OxmlElement('w:rFonts')
            rFonts.set(qn('w:ascii'), font_name)
            rFonts.set(qn('w:hAnsi'), font_name)
            rFonts.set(qn('w:cs'), font_name)
            rPr.append(rFonts)

def set_repeat_table_header(row):
    """คำสั่ง XML ให้แถวหัวตารางแสดงซ้ำที่ด้านบนทุกหน้าอัตโนมัติเมื่อขึ้นหน้าใหม่"""
    trPr = row._tr.get_or_add_trPr()
    tblHeader = OxmlElement('w:tblHeader')
    trPr.append(tblHeader)

def process_doc_placeholders(doc, replacements):
    """แทนที่ตัวแปร พร้อมสร้างฟิลด์เลขหน้าอัตโนมัติเมื่อพบ {{ page_no }}"""
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if "{{ page_no }}" in cell.text or "{{page_no}}" in cell.text:
                    for p in cell.paragraphs:
                        if "{{ page_no }}" in p.text or "{{page_no}}" in p.text:
                            p.text = p.text.replace("{{ page_no }}", "").replace("{{page_no}}", "")
                            run = p.add_run()
                            add_page_number_field(run)
                            run.font.name = "TH SarabunPSK"
                            run.font.size = Pt(14)
                for k, v in replacements.items():
                    if k in cell.text:
                        for p in cell.paragraphs:
                            if k in p.text:
                                p.text = p.text.replace(k, str(v))

    for p in doc.paragraphs:
        if "{{ page_no }}" in p.text or "{{page_no}}" in p.text:
            p.text = p.text.replace("{{ page_no }}", "").replace("{{page_no}}", "")
            run = p.add_run()
            add_page_number_field(run)
            run.font.name = "TH SarabunPSK"
            run.font.size = Pt(14)
        for k, v in replacements.items():
            if k in p.text:
                p.text = p.text.replace(k, str(v))

# ----------------------------------------------------
# 3. ฟังก์ชัน AI สกัดเนื้อหา (Multi-Model Dynamic Fallback หนี 503)
# ----------------------------------------------------
def get_file_content_for_ai(file_bytes: bytes, file_name: str, mime_type: str):
    if file_name.endswith(".docx"):
        try:
            doc = Document(io.BytesIO(file_bytes))
            text_lines = []
            for p in doc.paragraphs:
                if p.text.strip():
                    text_lines.append(p.text.strip())
            for tbl in doc.tables:
                for row in tbl.rows:
                    row_data = [c.text.strip().replace("\n", " ") for c in row.cells if c.text.strip()]
                    if row_data:
                        text_lines.append(" | ".join(row_data))
            return "\n".join(text_lines)
        except Exception:
            pass
    return types.Part.from_bytes(data=file_bytes, mime_type=mime_type)

def extract_course_plan(api_key: str, content_data, total_weeks: int, course_name: str):
    client = genai.Client(api_key=api_key)
    prompt = f"""
คุณคือผู้เชี่ยวชาญการจัดทำหลักสูตรและโครงการสอนระดับอาชีวศึกษา (มาตรฐาน สอศ.)
จงอ่านข้อมูลจากตารางวิเคราะห์งาน/หลักสูตรวิชา '{course_name}' แล้วสังเคราะห์จัดทำเป็น 'ตารางโครงการสอนรายสัปดาห์' ให้ครบถ้วนจำนวน {total_weeks} สัปดาห์ (สัปดาห์ที่ 1 ถึง {total_weeks})

*** ข้อกำหนดข้อมูลในแต่ละสัปดาห์ (เน้นความเชื่อมโยง 100%): ***
1. week: ตัวเลขสัปดาห์ (1 ถึง {total_weeks})
2. topic: ชื่อหน่วยและเรื่อง (บรรทัดแรก: หน่วยที่... ชื่อหน่วย / บรรทัดสอง: เรื่อง...)
3. tp: จุดประสงค์เชิงพฤติกรรม 2-3 ข้อ (สังเคราะห์จากสมรรถนะย่อย ความรู้ และทักษะ)
4. act: ขั้นตอน Active Learning 4 ขั้นตอน (1.ขั้นนำเข้าสู่บทเรียน 2.ขั้นให้ความรู้/ศึกษาค้นคว้า 3.ขั้นฝึกปฏิบัติการ 4.ขั้นสรุปและประเมินผล)
5. media: สื่อการเรียนรู้ที่ **สอดคล้องกับสิ่งที่นักเรียนใช้ในกิจกรรมปฏิบัติจริง** (เช่น แบบฟอร์มใบสั่งซื้อ PO, ซอฟต์แวร์ ERP จำลอง, แผนผัง Layout คลังสินค้า พร้อมสไลด์และใบงาน)
6. assess: การวัดผลที่ **สอดคล้องกับกิจกรรมจริง** (เช่น แบบประเมินทักษะ Rubric, แบบทดสอบย่อย, ตรวจผลงานใบงาน)

ตอบกลับเป็น Pure JSON Array เท่านั้น ห้ามใส่ markdown code block ครอบ:
[
  {{
    "week": 1,
    "topic": "หน่วยที่ 1 ความรู้เบื้องต้นเกี่ยวกับซัพพลายเชน\\nเรื่อง ความหมายและขอบเขต",
    "tp": "- อธิบายความหมายและขอบเขตได้\\n- เปรียบเทียบความแตกต่างได้",
    "act": "1. ขั้นนำ: เปิดคลิปวิดีโอ...\\n2. ขั้นสอน: บรรยายสไลด์...\\n3. ขั้นปฏิบัติ: ทำใบงานวิเคราะห์...\\n4. ขั้นสรุป: สรุปร่วมกัน",
    "media": "- สไลด์มัลติมีเดีย หน่วยที่ 1\\n- คลิปวิดีโอวงจรชีวิตผลิตภัณฑ์\\n- ใบงานที่ 1.1 เรื่อง โครงสร้างซัพพลายเชน",
    "assess": "- แบบประเมินใบงานที่ 1.1\\n- แบบทดสอบย่อยท้ายคาบ\\n- แบบสังเกตพฤติกรรมการทำงาน"
  }}
]
"""
    if isinstance(content_data, str):
        full_contents = f"ข้อมูลตารางวิเคราะห์งาน:\n{content_data[:12000]}\n\n{prompt}"
    else:
        full_contents = [content_data, prompt]

    # กระจายโมเดลสำรองเพื่อเลี่ยงเซิร์ฟเวอร์เต็ม (503 UNAVAILABLE)
    candidate_models = ["gemini-2.5-flash", "gemini-2.5-pro", "gemini-3.6-flash"]
    last_error = None

    for model_name in candidate_models:
        for attempt in range(2):
            try:
                response = client.models.generate_content(
                    model=model_name,
                    contents=full_contents,
                    config=types.GenerateContentConfig(response_mime_type="application/json")
                )
                return json.loads(response.text)
            except Exception as e:
                last_error = e
                err_msg = str(e)
                if "503" in err_msg or "UNAVAILABLE" in err_msg:
                    time.sleep(3)
                    continue
                else:
                    break

    raise last_error

# ----------------------------------------------------
# 4. ส่วนรับข้อมูลหน้าเว็บ
# ----------------------------------------------------
with st.sidebar:
    st.markdown("### ⚙️ การตั้งค่าระบบ")
    st.success("🟢 STATUS: AUTHORIZED")
    api_key = st.text_input("🔑 Gemini API Key:", type="password", placeholder="AIzaSy...")
    st.markdown("[👉 รับ API Key จาก Google AI Studio](https://aistudio.google.com/)")
    st.markdown("---")
    teacher_name = st.text_input("ชื่อ-สกุล ครูผู้สอน:", value="นายณัฐวุฒิ หล้าปงสาย")
    dept_name = st.text_input("แผนกวิชา:", value="การจัดการโลจิสติกส์และซัพพลายเชน")
    if st.button("🔒 ล็อกระบบกลับ"):
        st.session_state.authenticated = False
        st.rerun()

st.markdown("""
<div class="hero-banner">
    <div class="hero-title">📋 ระบบจัดทำโครงการสอนอัจฉริยะ (สอศ.)</div>
    <div class="hero-desc">
        สกัดตารางวิเคราะห์งานสู่โครงการสอนอัตโนมัติ รองรับ ปวส. (15 สัปดาห์ | ปี 1-2) และ ปวช. (18 สัปดาห์ | ปี 1-3)<br>
        ระบบออกแบบกิจกรรมเชิงรุก สื่อ และการวัดผลที่เชื่อมโยงตรงตามบริบท พร้อมซ้ำหัวตารางทุกหน้า
    </div>
</div>
""", unsafe_allow_html=True)

col_file, col_info = st.columns([1, 1], gap="large")

with col_file:
    st.markdown('<div class="box-header">📁 1. เอกสารนำเข้า</div>', unsafe_allow_html=True)
    template_file = st.file_uploader("แบบฟอร์มวิทยาลัย (templet.docx):", type=["docx"])
    analysis_file = st.file_uploader("ไฟล์ตารางวิเคราะห์งาน (docx/pdf):", type=["docx", "pdf"])

with col_info:
    st.markdown('<div class="box-header">🎯 2. ข้อมูลวิชาและระดับชั้น</div>', unsafe_allow_html=True)
    degree_select = st.selectbox("ระดับคุณวุฒิการศึกษา:", ["ประกาศนียบัตรวิชาชีพชั้นสูง (ปวส.)", "ประกาศนียบัตรวิชาชีพ (ปวช.)"])
    is_pvs = "ปวส." in degree_select
    
    c_y, c_w, c_h = st.columns(3)
    with c_y:
        year_opts = [1, 2] if is_pvs else [1, 2, 3]
        year_input = st.selectbox("ระดับชั้นปี:", year_opts, index=0)
    with c_w:
        default_weeks = 15 if is_pvs else 18
        weeks_input = st.number_input("สัปดาห์ต่อภาคเรียน:", min_value=1, max_value=22, value=default_weeks)
    with c_h:
        hours_input = st.number_input("ชั่วโมงต่อสัปดาห์:", min_value=1, max_value=10, value=4)
        
    c_code, c_sem = st.columns(2)
    with c_code:
        course_code_input = st.text_input("รหัสวิชา:", value="31401-2007")
    with c_sem:
        sem_input = st.text_input("ภาคเรียนที่:", value="1/2569")
        
    course_name_input = st.text_input("ชื่อวิชา:", value="การจัดการโลจิสติกส์และซัพพลายเชน")

st.markdown("<br>", unsafe_allow_html=True)

# ----------------------------------------------------
# 5. ประมวลผลและสร้างไฟล์ Word
# ----------------------------------------------------
if st.button("🚀 ประมวลผลและสร้างโครงการสอน (Generate Word)", type="primary", use_container_width=True):
    if not api_key:
        st.warning("⚠️ กรุณาระบุ Gemini API Key ในแถบด้านซ้าย")
    elif not template_file or not analysis_file:
        st.warning("⚠️ กรุณาแนบทั้ง 'แบบฟอร์มวิทยาลัย' และ 'ไฟล์ตารางวิเคราะห์งาน'")
    else:
        with st.status("⚡ กำลังสร้างโครงการสอนมาตรฐาน สอศ. ...", expanded=True) as status:
            try:
                st.write("📖 กำลังเตรียมข้อมูลจากตารางวิเคราะห์งาน...")
                raw_bytes = analysis_file.getvalue()
                mime = "application/pdf" if analysis_file.name.endswith(".pdf") else "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
                content_data = get_file_content_for_ai(raw_bytes, analysis_file.name, mime)

                st.write("🤖 กำลังวิเคราะห์เนื้อหา เชื่อมโยงกิจกรรม สื่อ และการวัดผลเฉพาะบริบท...")
                plans = extract_course_plan(
                    api_key=api_key,
                    content_data=content_data,
                    total_weeks=weeks_input,
                    course_name=course_name_input
                )
                
                st.write("📝 กำลังบรรจุข้อมูลลงแบบฟอร์ม Word และจัดระบบเลขหน้า...")
                doc = Document(template_file)
                
                CHECK, UNCHECK = "☑", "☐"
                if is_pvs:
                    deg_full = "ประกาศนียบัตรวิชาชีพชั้นสูง (ปวส.)"
                    y_disp = f"{CHECK if year_input == 1 else UNCHECK} ปี 1   {CHECK if year_input == 2 else UNCHECK} ปี 2"
                else:
                    deg_full = "ประกาศนียบัตรวิชาชีพ (ปวช.)"
                    y_disp = f"{CHECK if year_input == 1 else UNCHECK} ปี 1   {CHECK if year_input == 2 else UNCHECK} ปี 2   {CHECK if year_input == 3 else UNCHECK} ปี 3"

                replacements = {
                    "{{ degree_title }}": deg_full,
                    "{{ course_code }}": course_code_input,
                    "{{ course_name }}": course_name_input,
                    "{{ hours_per_week }}": str(hours_input),
                    "{{ total_weeks }}": str(weeks_input),
                    "{{ semester }}": sem_input,
                    "{{ year_display }}": y_disp,
                    "{{ sheet_no }}": "1",
                    "{{sheet_no}}": "1",
                    "{{ cb_y1 }}": CHECK if year_input == 1 else UNCHECK,
                    "{{ cb_y2 }}": CHECK if year_input == 2 else UNCHECK,
                    "{{ cb_y3 }}": CHECK if year_input == 3 else UNCHECK
                }
                process_doc_placeholders(doc, replacements)
                
                target_table = None
                header_row_index = -1
                for tbl in doc.tables:
                    for idx, row in enumerate(tbl.rows):
                        row_text = " ".join([c.text for c in row.cells])
                        if "ส.ป." in row_text or "Teaching Point" in row_text or "กิจกรรม" in row_text:
                            target_table = tbl
                            header_row_index = idx
                            break
                    if target_table:
                        break
                if not target_table:
                    target_table = doc.tables[0]
                    header_row_index = 0
                
                # ซ้ำแถวหัวตารางทุกหน้าเมื่อขึ้นหน้าใหม่
                if header_row_index >= 0:
                    set_repeat_table_header(target_table.rows[header_row_index])
                
                # หยอดข้อมูล 6 คอลัมน์แบบตรงช่อง 100%
                for item in plans:
                    new_row = target_table.add_row()
                    row_cells = new_row.cells
                    num_cols = len(row_cells)
                    
                    if num_cols >= 6:
                        # ช่อง 0: ส.ป.
                        row_cells[0].text = str(item.get("week", ""))
                        row_cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
                        
                        # ช่องขวาสุด: วัดผล
                        row_cells[-1].text = str(item.get("assess", ""))
                        # ช่องรองสุดท้าย: สื่อ
                        row_cells[-2].text = str(item.get("media", ""))
                        # ช่องกิจกรรม
                        row_cells[-3].text = str(item.get("act", ""))
                        # ช่อง Teaching Point
                        row_cells[-4].text = str(item.get("tp", ""))
                        
                        # ช่องหัวข้อ (เติมเต็มช่องตรงกลางระหว่าง ส.ป. และ Teaching Point)
                        for c_idx in range(1, num_cols - 4):
                            row_cells[c_idx].text = str(item.get("topic", ""))
                    else:
                        row_cells[0].text = str(item.get("week", ""))
                    
                    for cell in row_cells:
                        set_cell_font(cell, font_name="TH SarabunPSK", font_size=Pt(14))
                
                out_stream = io.BytesIO()
                doc.save(out_stream)
                out_stream.seek(0)
                
                status.update(label="✅ ดำเนินการสร้างโครงการสอนสำเร็จ!", state="complete", expanded=False)
                st.balloons()
                st.success("🎉 ระบบจัดทำโครงการสอนเสร็จสมบูรณ์เรียบร้อยแล้ว")
                
                st.download_button(
                    label=f"📥 ดาวน์โหลดโครงการสอน_{course_code_input}.docx",
                    data=out_stream,
                    file_name=f"โครงการสอน_{course_code_input}_{sem_input.replace('/', '_')}.docx",
                    mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    use_container_width=True
                )
            except Exception as err:
                status.update(label="❌ เกิดข้อผิดพลาด", state="error")
                st.error(f"รายละเอียดข้อผิดพลาด: {str(err)}")
